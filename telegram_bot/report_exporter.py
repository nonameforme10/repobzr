"""
Bazar (SalesTrack) — Report Exporter Service (Python)
Generates high-resolution PNG images via Pillow (PIL)
and executive Excel (.xlsx) workbooks via openpyxl.

STRICT ARCHITECTURAL RULE:
Exporters are pure rendering functions. They NEVER query the database.
They consume the normalized reportData structure produced by reportService.
"""

from io import BytesIO
import os
from typing import Dict, Any, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageDraw, ImageFont


def format_soom(n: Any) -> str:
    """Format integer into space-separated string (e.g. 1400000 -> '1 400 000')."""
    if n is None:
        return "0"
    try:
        val = round(float(n))
        return f"{val:,}".replace(",", " ")
    except Exception:
        return "0"


def _get_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    """Load system font with graceful fallbacks."""
    font_names = ["arialbd.ttf" if bold else "arial.ttf", "segoeuib.ttf" if bold else "segoeui.ttf", "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"]
    for name in font_names:
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            continue
    return ImageFont.load_default()


# ============================================================================
# 1. SINGLE-DAY EXCEL EXPORT (.xlsx)
# ============================================================================
def generate_excel_report(report_data: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kunlik hisobot"
    ws.views.sheetView[0].showGridLines = True

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 22

    # Styles
    title_font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    date_font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    date_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    meta_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    meta_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    # 1. Title Banner
    ws.merge_cells("A1:E1")
    cell_a1 = ws["A1"]
    cell_a1.value = "BAZAR — KUNLIK SAVDO HISOBOTI"
    cell_a1.font = title_font
    cell_a1.fill = title_fill
    cell_a1.alignment = center_align
    ws.row_dimensions[1].height = 36

    # 2. Date Banner
    ws.merge_cells("A2:E2")
    cell_a2 = ws["A2"]
    cell_a2.value = f"HISOBOT SANASI: {report_data['date']}"
    cell_a2.font = date_font
    cell_a2.fill = date_fill
    cell_a2.alignment = center_align
    ws.row_dimensions[2].height = 28

    # 3. Metadata
    ws.merge_cells("A3:E3")
    cell_a3 = ws["A3"]
    cell_a3.value = "Valyuta: UZS (soʻm)   |   Vaqt mintaqasi: Asia/Tashkent (UTC+05:00)"
    cell_a3.font = meta_font
    cell_a3.fill = meta_fill
    cell_a3.alignment = center_align
    ws.row_dimensions[3].height = 20

    # 4. KPI Summary Row
    ws.row_dimensions[4].height = 10  # gap

    kpi_headers = ["JAMI TUSHUM", "JAMI SOTILDI", "CHEGIRMA", "CHIQIM", "SOF KASSA"]
    cols = ["A", "B", "C", "D", "E"]
    for c_letter, h in zip(cols, kpi_headers):
        cell = ws[f"{c_letter}5"]
        cell.value = h
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="475569")
        cell.alignment = center_align

    totals = report_data["totals"]
    ws["A6"] = f"{format_soom(totals['netSales'])} so'm"
    ws["A6"].font = Font(name="Segoe UI", size=12, bold=True, color="16A34A")

    ws["B6"] = f"{totals['itemsSold']} dona ({totals['transactionCount']} ta savdo)"
    ws["B6"].font = Font(name="Segoe UI", size=11, bold=True)

    ws["C6"] = f"{format_soom(totals['discount'])} so'm"
    ws["C6"].font = Font(name="Segoe UI", size=11, bold=True)

    ws["D6"] = f"−{format_soom(totals['chiqim'])} so'm"
    ws["D6"].font = Font(name="Segoe UI", size=11, bold=True, color="DC2626")

    ws["E6"] = f"{format_soom(totals['netCash'])} so'm"
    ws["E6"].font = Font(name="Segoe UI", size=12, bold=True, color="0284C7")

    for c_letter in cols:
        ws[f"{c_letter}6"].alignment = center_align
    ws.row_dimensions[6].height = 28

    curr_row = 8

    # 5. Chiqim Section
    cash_outs = report_data.get("cashOuts", [])
    if cash_outs:
        ws.merge_cells(f"A{curr_row}:E{curr_row}")
        c_title = ws[f"A{curr_row}"]
        c_title.value = "💸 KASSADAN CHIQIMLAR"
        c_title.font = Font(name="Segoe UI", size=11, bold=True, color="DC2626")
        c_title.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        headers = ["#", "Sabab / Izoh", "Vaqt", "", "Miqdor"]
        for c_letter, h in zip(cols, headers):
            cell = ws[f"{c_letter}{curr_row}"]
            cell.value = h
            cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
            cell.alignment = center_align
        curr_row += 1

        for idx, c in enumerate(cash_outs, start=1):
            ws[f"A{curr_row}"] = idx
            ws[f"A{curr_row}"].alignment = center_align

            ws[f"B{curr_row}"] = c.get("reason", "Kassadan chiqim")
            ws[f"B{curr_row}"].alignment = left_align

            ws[f"C{curr_row}"] = c.get("timeFormatted", "")
            ws[f"C{curr_row}"].alignment = center_align

            ws[f"E{curr_row}"] = f"−{format_soom(c['amount'])} so'm"
            ws[f"E{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
            ws[f"E{curr_row}"].alignment = right_align
            curr_row += 1

        curr_row += 2

    # 6. Items Sold Section
    ws.merge_cells(f"A{curr_row}:E{curr_row}")
    items_title = ws[f"A{curr_row}"]
    items_title.value = "📦 SOTILGAN MAHSULOTLAR"
    items_title.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    items_title.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws.row_dimensions[curr_row].height = 24
    curr_row += 1

    tbl_headers = ["#", "Mahsulot nomi", "Soni", "Dona narxi", "Jami summa"]
    for c_letter, h in zip(cols, tbl_headers):
        cell = ws[f"{c_letter}{curr_row}"]
        cell.value = h
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.alignment = center_align
    curr_row += 1

    items = report_data.get("items", [])
    if not items:
        ws.merge_cells(f"A{curr_row}:E{curr_row}")
        ws[f"A{curr_row}"] = "Bu kunda sotuvlar qayd etilmagan"
        ws[f"A{curr_row}"].alignment = center_align
        curr_row += 1
    else:
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        for idx, item in enumerate(items, start=1):
            ws[f"A{curr_row}"] = idx
            ws[f"A{curr_row}"].alignment = center_align

            ws[f"B{curr_row}"] = item["name"]
            ws[f"B{curr_row}"].alignment = left_align

            ws[f"C{curr_row}"] = item["quantity"]
            ws[f"C{curr_row}"].alignment = center_align

            ws[f"D{curr_row}"] = f"{format_soom(item['unitPrice'])} so'm"
            ws[f"D{curr_row}"].alignment = right_align

            ws[f"E{curr_row}"] = f"{format_soom(item['total'])} so'm"
            ws[f"E{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
            ws[f"E{curr_row}"].alignment = right_align

            if idx % 2 == 0:
                for c_letter in cols:
                    ws[f"{c_letter}{curr_row}"].fill = zebra_fill
            curr_row += 1

        # Totals Row
        ws.merge_cells(f"A{curr_row}:B{curr_row}")
        ws[f"A{curr_row}"] = "JAMI"
        ws[f"A{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
        ws[f"A{curr_row}"].alignment = right_align

        ws[f"C{curr_row}"] = totals["itemsSold"]
        ws[f"C{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
        ws[f"C{curr_row}"].alignment = center_align

        ws[f"E{curr_row}"] = f"{format_soom(totals['grossSales'])} so'm"
        ws[f"E{curr_row}"].font = Font(name="Segoe UI", size=11, bold=True, color="16A34A")
        ws[f"E{curr_row}"].alignment = right_align

        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        for c_letter in cols:
            ws[f"{c_letter}{curr_row}"].fill = total_fill

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ============================================================================
# 2. 7-DAY EXCEL EXPORT (.xlsx)
# ============================================================================
def generate_7day_excel_report(seven_day_data: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "7 kunlik hisobot"
    ws.views.sheetView[0].showGridLines = True

    cols = ["A", "B", "C", "D", "E", "F", "G"]
    widths = [14, 20, 16, 22, 14, 18, 22]
    for c_letter, w in zip(cols, widths):
        ws.column_dimensions[c_letter].width = w

    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # 1. Title
    ws.merge_cells("A1:G1")
    ws["A1"] = "BAZAR — 7 KUNLIK SAVDO TAHLILI"
    ws["A1"].font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    ws["A1"].alignment = center_align
    ws.row_dimensions[1].height = 36

    # 2. Range
    ws.merge_cells("A2:G2")
    ws["A2"] = f"DAVR: {seven_day_data['startDate']} dan {seven_day_data['endDate']} gacha"
    ws["A2"].font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    ws["A2"].fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    ws["A2"].alignment = center_align
    ws.row_dimensions[2].height = 28

    # 3. Metadata
    ws.merge_cells("A3:G3")
    ws["A3"] = "Valyuta: UZS (soʻm)   |   Asia/Tashkent (UTC+05:00)"
    ws["A3"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    ws["A3"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws["A3"].alignment = center_align

    ws.row_dimensions[4].height = 12

    # Headers
    headers = ["Sana", "Savdo (Gross)", "Chegirma", "Sof Savdo (Net)", "Sotildi", "Chiqim", "Sof Kassa"]
    for c_letter, h in zip(cols, headers):
        cell = ws[f"{c_letter}5"]
        cell.value = h
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.alignment = center_align
    ws.row_dimensions[5].height = 26

    curr_row = 6
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, d in enumerate(seven_day_data["dailyReports"], start=1):
        t = d["totals"]
        ws[f"A{curr_row}"] = d.get("displayDate") or d["date"]
        ws[f"A{curr_row}"].alignment = center_align

        ws[f"B{curr_row}"] = f"{format_soom(t['grossSales'])} so'm"
        ws[f"B{curr_row}"].alignment = right_align

        ws[f"C{curr_row}"] = f"{format_soom(t['discount'])} so'm"
        ws[f"C{curr_row}"].alignment = right_align

        ws[f"D{curr_row}"] = f"{format_soom(t['netSales'])} so'm"
        ws[f"D{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
        ws[f"D{curr_row}"].alignment = right_align

        ws[f"E{curr_row}"] = f"{t['itemsSold']} dona"
        ws[f"E{curr_row}"].alignment = center_align

        ws[f"F{curr_row}"] = f"−{format_soom(t['chiqim'])} so'm" if t["chiqim"] > 0 else "0"
        ws[f"F{curr_row}"].font = Font(name="Segoe UI", size=10, color="DC2626" if t["chiqim"] > 0 else "64748B")
        ws[f"F{curr_row}"].alignment = right_align

        ws[f"G{curr_row}"] = f"{format_soom(t['netCash'])} so'm"
        ws[f"G{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True, color="0284C7")
        ws[f"G{curr_row}"].alignment = right_align

        if idx % 2 == 0:
            for c_letter in cols:
                ws[f"{c_letter}{curr_row}"].fill = zebra_fill
        curr_row += 1

    # Rollup Row
    rollup = seven_day_data["rollup"]
    ws[f"A{curr_row}"] = "JAMI (7 KUN)"
    ws[f"A{curr_row}"].font = Font(name="Segoe UI", size=11, bold=True)
    ws[f"A{curr_row}"].alignment = center_align

    ws[f"B{curr_row}"] = f"{format_soom(rollup['grossSales'])} so'm"
    ws[f"B{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
    ws[f"B{curr_row}"].alignment = right_align

    ws[f"C{curr_row}"] = f"{format_soom(rollup['discount'])} so'm"
    ws[f"C{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
    ws[f"C{curr_row}"].alignment = right_align

    ws[f"D{curr_row}"] = f"{format_soom(rollup['netSales'])} so'm"
    ws[f"D{curr_row}"].font = Font(name="Segoe UI", size=11, bold=True, color="16A34A")
    ws[f"D{curr_row}"].alignment = right_align

    ws[f"E{curr_row}"] = f"{rollup['itemsSold']} dona"
    ws[f"E{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True)
    ws[f"E{curr_row}"].alignment = center_align

    ws[f"F{curr_row}"] = f"−{format_soom(rollup['chiqim'])} so'm"
    ws[f"F{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    ws[f"F{curr_row}"].alignment = right_align

    ws[f"G{curr_row}"] = f"{format_soom(rollup['netCash'])} so'm"
    ws[f"G{curr_row}"].font = Font(name="Segoe UI", size=11, bold=True, color="0284C7")
    ws[f"G{curr_row}"].alignment = right_align

    rollup_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    for c_letter in cols:
        ws[f"{c_letter}{curr_row}"].fill = rollup_fill
    ws.row_dimensions[curr_row].height = 28

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ============================================================================
# 3. SINGLE-DAY PNG EXPORT (860px Pillow Image)
# ============================================================================
def generate_png_report(report_data: Dict[str, Any]) -> bytes:
    width = 860
    padding = 28

    items = report_data.get("items", [])
    cash_outs = report_data.get("cashOuts", [])

    items_count = len(items)
    cash_out_count = len(cash_outs)

    height = 340
    if cash_out_count > 0:
        height += 60 + (cash_out_count * 36) + 16
    if items_count > 0:
        height += 60 + (min(items_count, 25) * 36) + 50
    else:
        height += 80

    image = Image.new("RGB", (width, height), color=(10, 13, 20))
    draw = ImageDraw.Draw(image)

    # Card background
    draw.rounded_rectangle(
        [(padding, padding), (width - padding, height - padding)],
        radius=16,
        fill=(20, 26, 38),
        outline=(35, 45, 63),
        width=2
    )

    y = padding + 24

    # Header Brand
    draw.text((padding + 24, y), "BAZAR", fill=(212, 175, 55), font=_get_font(24, bold=True))
    draw.text((padding + 24, y + 32), "Kunlik savdo va kassa hisoboti", fill=(148, 163, 184), font=_get_font(13))

    # Date Badge
    date_text = f"📅 {report_data['date']}"
    date_font = _get_font(13, bold=True)
    bbox = date_font.getbbox(date_text)
    badge_w = (bbox[2] - bbox[0]) + 24
    badge_x = width - padding - 24 - badge_w

    draw.rounded_rectangle(
        [(badge_x, y + 4), (badge_x + badge_w, y + 38)],
        radius=17,
        fill=(30, 41, 59)
    )
    draw.text((badge_x + 12, y + 12), date_text, fill=(248, 250, 252), font=date_font)

    y += 70

    # Divider
    draw.line([(padding + 24, y), (width - padding - 24, y)], fill=(35, 45, 63), width=1)
    y += 20

    # 4 KPI Cards
    kpi_w = (width - padding * 2 - 48 - (3 * 12)) // 4
    kpi_h = 90
    totals = report_data["totals"]

    kpi_defs = [
        {
            "title": "JAMI TUSHUM",
            "val": f"{format_soom(totals['netSales'])} so'm",
            "sub": f"{totals['transactionCount']} ta savdo",
            "bg": (20, 45, 30),
            "border": (34, 197, 94),
            "val_color": (34, 197, 94)
        },
        {
            "title": "JAMI SOTILDI",
            "val": f"{totals['itemsSold']} dona",
            "sub": f"Chegirma: −{format_soom(totals['discount'])}" if totals["discount"] > 0 else "Mahsulotlar soni",
            "bg": (26, 35, 51),
            "border": (42, 55, 74),
            "val_color": (248, 250, 252)
        },
        {
            "title": "KASSADAN CHIQIM",
            "val": f"−{format_soom(totals['chiqim'])} so'm",
            "sub": f"{cash_out_count} ta chiqim",
            "bg": (45, 20, 25),
            "border": (239, 68, 68),
            "val_color": (239, 68, 68)
        },
        {
            "title": "SOF KASSA",
            "val": f"{format_soom(totals['netCash'])} so'm",
            "sub": "Kassadagi naqd pul",
            "bg": (15, 35, 55),
            "border": (56, 189, 248),
            "val_color": (56, 189, 248)
        }
    ]

    for idx, k in enumerate(kpi_defs):
        kx = padding + 24 + idx * (kpi_w + 12)
        draw.rounded_rectangle(
            [(kx, y), (kx + kpi_w, y + kpi_h)],
            radius=10,
            fill=k["bg"],
            outline=k["border"],
            width=1
        )
        draw.text((kx + 12, y + 12), k["title"], fill=(148, 163, 184), font=_get_font(10, bold=True))
        draw.text((kx + 12, y + 36), k["val"], fill=k["val_color"], font=_get_font(15, bold=True))
        draw.text((kx + 12, y + 64), k["sub"], fill=(100, 116, 139), font=_get_font(11))

    y += kpi_h + 24

    # Chiqim Section
    if cash_out_count > 0:
        draw.text((padding + 24, y), "💸 KASSADAN CHIQIMLAR", fill=(239, 68, 68), font=_get_font(12, bold=True))
        y += 20
        table_w = width - padding * 2 - 48
        for idx, c in enumerate(cash_outs):
            row_bg = (24, 32, 48) if idx % 2 == 0 else (20, 26, 38)
            draw.rectangle([(padding + 24, y), (padding + 24 + table_w, y + 28)], fill=row_bg)
            draw.text((padding + 34, y + 6), c.get("timeFormatted", ""), fill=(148, 163, 184), font=_get_font(11))
            draw.text((padding + 90, y + 6), c.get("reason", "Kassadan chiqim"), fill=(248, 250, 252), font=_get_font(12, bold=True))
            amt_str = f"−{format_soom(c['amount'])} so'm"
            draw.text((width - padding - 36 - 130, y + 6), amt_str, fill=(220, 38, 38), font=_get_font(12, bold=True))
            y += 30
        y += 16

    # Items Section
    draw.text((padding + 24, y), "📦 SOTILGAN MAHSULOTLAR", fill=(248, 250, 252), font=_get_font(13, bold=True))
    y += 22

    table_x = padding + 24
    table_w = width - padding * 2 - 48

    # Table Header
    draw.rectangle([(table_x, y), (table_x + table_w, y + 28)], fill=(30, 41, 59))
    draw.text((table_x + 10, y + 6), "#", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + 40, y + 6), "Mahsulot nomi", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + table_w - 240, y + 6), "Soni", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + table_w - 160, y + 6), "Narxi", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + table_w - 70, y + 6), "Jami", fill=(148, 163, 184), font=_get_font(11, bold=True))
    y += 28

    if items_count == 0:
        draw.text((table_x + 20, y + 10), "Bu kunda sotuvlar qayd etilmagan", fill=(100, 116, 139), font=_get_font(12))
    else:
        for idx, it in enumerate(items[:25], start=1):
            row_bg = (24, 32, 48) if idx % 2 == 1 else (20, 26, 38)
            draw.rectangle([(table_x, y), (table_x + table_w, y + 30)], fill=row_bg)

            draw.text((table_x + 10, y + 7), str(idx), fill=(100, 116, 139), font=_get_font(11))

            name = it["name"]
            if len(name) > 34:
                name = name[:32] + "…"
            draw.text((table_x + 40, y + 7), name, fill=(248, 250, 252), font=_get_font(12, bold=True))

            draw.text((table_x + table_w - 240, y + 7), f"{it['quantity']} dona", fill=(226, 232, 240), font=_get_font(12))
            draw.text((table_x + table_w - 160, y + 7), format_soom(it["unitPrice"]), fill=(226, 232, 240), font=_get_font(12))
            draw.text((table_x + table_w - 70, y + 7), format_soom(it["total"]), fill=(34, 197, 94), font=_get_font(12, bold=True))
            y += 30

        if items_count > 25:
            draw.text(
                (table_x + 10, y + 8),
                f"... va yana {items_count - 25} ta mahsulot (to'liq ro'yxat Excel faylida)",
                fill=(100, 116, 139),
                font=_get_font(11)
            )

    bio = BytesIO()
    image.save(bio, format="PNG")
    return bio.getvalue()


# ============================================================================
# 4. 7-DAY PNG EXPORT (860px Pillow Image)
# ============================================================================
def generate_7day_png_report(seven_day_data: Dict[str, Any]) -> bytes:
    width = 860
    padding = 28
    row_height = 34
    height = 300 + (len(seven_day_data["dailyReports"]) * row_height) + 80

    image = Image.new("RGB", (width, height), color=(10, 13, 20))
    draw = ImageDraw.Draw(image)

    # Card background
    draw.rounded_rectangle(
        [(padding, padding), (width - padding, height - padding)],
        radius=16,
        fill=(20, 26, 38),
        outline=(35, 45, 63),
        width=2
    )

    y = padding + 24

    # Header
    draw.text((padding + 24, y), "BAZAR", fill=(212, 175, 55), font=_get_font(24, bold=True))
    draw.text((padding + 24, y + 32), "7 kunlik savdo tahlili", fill=(148, 163, 184), font=_get_font(13))

    # Range Badge
    range_text = f"📅 {seven_day_data['startDate']} — {seven_day_data['endDate']}"
    r_font = _get_font(12, bold=True)
    bbox = r_font.getbbox(range_text)
    badge_w = (bbox[2] - bbox[0]) + 24
    badge_x = width - padding - 24 - badge_w

    draw.rounded_rectangle(
        [(badge_x, y + 4), (badge_x + badge_w, y + 38)],
        radius=17,
        fill=(30, 41, 59)
    )
    draw.text((badge_x + 12, y + 12), range_text, fill=(248, 250, 252), font=r_font)

    y += 70

    # Divider
    draw.line([(padding + 24, y), (width - padding - 24, y)], fill=(35, 45, 63), width=1)
    y += 20

    # 4 Rollup KPI Cards
    kpi_w = (width - padding * 2 - 48 - (3 * 12)) // 4
    rollup = seven_day_data["rollup"]

    kpi_defs = [
        {"title": "7 KUNLIK SAVDO", "val": f"{format_soom(rollup['netSales'])} so'm", "color": (34, 197, 94)},
        {"title": "JAMI SOTILDI", "val": f"{rollup['itemsSold']} dona", "color": (248, 250, 252)},
        {"title": "JAMI CHIQIM", "val": f"−{format_soom(rollup['chiqim'])} so'm", "color": (239, 68, 68)},
        {"title": "SOF KASSA", "val": f"{format_soom(rollup['netCash'])} so'm", "color": (56, 189, 248)}
    ]

    for idx, k in enumerate(kpi_defs):
        kx = padding + 24 + idx * (kpi_w + 12)
        draw.rounded_rectangle(
            [(kx, y), (kx + kpi_w, y + 75)],
            radius=10,
            fill=(26, 35, 51),
            outline=(42, 55, 74),
            width=1
        )
        draw.text((kx + 12, y + 10), k["title"], fill=(148, 163, 184), font=_get_font(10, bold=True))
        draw.text((kx + 12, y + 36), k["val"], fill=k["color"], font=_get_font(15, bold=True))

    y += 95

    # Table Header
    table_x = padding + 24
    table_w = width - padding * 2 - 48

    draw.rectangle([(table_x, y), (table_x + table_w, y + 30)], fill=(30, 41, 59))
    draw.text((table_x + 16, y + 8), "Sana", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + 160, y + 8), "Savdo (Net)", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + 320, y + 8), "Sotildi", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + 440, y + 8), "Chiqim", fill=(148, 163, 184), font=_get_font(11, bold=True))
    draw.text((table_x + table_w - 130, y + 8), "Sof Kassa", fill=(148, 163, 184), font=_get_font(11, bold=True))
    y += 30

    for idx, d in enumerate(seven_day_data["dailyReports"]):
        row_bg = (24, 32, 48) if idx % 2 == 1 else (20, 26, 38)
        draw.rectangle([(table_x, y), (table_x + table_w, y + row_height)], fill=row_bg)

        dt_label = d.get("displayDate") or d["date"]
        draw.text((table_x + 16, y + 8), dt_label, fill=(248, 250, 252), font=_get_font(12, bold=True))

        t = d["totals"]
        draw.text((table_x + 160, y + 8), f"{format_soom(t['netSales'])} so'm", fill=(34, 197, 94), font=_get_font(12))
        draw.text((table_x + 320, y + 8), f"{t['itemsSold']} dona", fill=(226, 232, 240), font=_get_font(12))

        chiqim_str = f"−{format_soom(t['chiqim'])}" if t["chiqim"] > 0 else "0"
        chiqim_col = (239, 68, 68) if t["chiqim"] > 0 else (100, 116, 139)
        draw.text((table_x + 440, y + 8), chiqim_str, fill=chiqim_col, font=_get_font(12))

        draw.text((table_x + table_w - 130, y + 8), f"{format_soom(t['netCash'])} so'm", fill=(56, 189, 248), font=_get_font(12, bold=True))
        y += row_height

    # Rollup Row
    draw.rectangle([(table_x, y), (table_x + table_w, y + 36)], fill=(30, 41, 59))
    draw.text((table_x + 16, y + 10), "JAMI (7 KUN)", fill=(212, 175, 55), font=_get_font(12, bold=True))
    draw.text((table_x + 160, y + 10), f"{format_soom(rollup['netSales'])} so'm", fill=(34, 197, 94), font=_get_font(12, bold=True))
    draw.text((table_x + 320, y + 10), f"{rollup['itemsSold']} dona", fill=(248, 250, 252), font=_get_font(12, bold=True))
    draw.text((table_x + 440, y + 10), f"−{format_soom(rollup['chiqim'])}", fill=(239, 68, 68), font=_get_font(12, bold=True))
    draw.text((table_x + table_w - 130, y + 10), f"{format_soom(rollup['netCash'])} so'm", fill=(56, 189, 248), font=_get_font(12, bold=True))

    bio = BytesIO()
    image.save(bio, format="PNG")
    return bio.getvalue()
